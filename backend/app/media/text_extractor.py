"""
Extract text from PDF and PPTX files, one entry per page/slide so that
source grounding ("Page 7") is preserved.
"""
from dataclasses import dataclass

from pptx import Presentation
from pypdf import PdfReader


@dataclass
class PageText:
    page_number: int  # 1-indexed
    text: str


def extract_pdf_pages(pdf_path: str) -> list[PageText]:
    try:
        reader = PdfReader(pdf_path)
    except Exception as e:
        raise RuntimeError(f"Could not read PDF: {e}")

    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(PageText(page_number=i, text=text))
    if not pages:
        raise RuntimeError("No extractable text found in PDF (it may be a scanned/image-only document).")
    return pages


def extract_pptx_slides(pptx_path: str) -> list[PageText]:
    try:
        prs = Presentation(pptx_path)
    except Exception as e:
        raise RuntimeError(f"Could not read PPTX: {e}")

    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        combined = "\n".join(t.strip() for t in texts if t.strip())
        if combined:
            slides.append(PageText(page_number=i, text=combined))
    if not slides:
        raise RuntimeError("No extractable text found in PPTX.")
    return slides
